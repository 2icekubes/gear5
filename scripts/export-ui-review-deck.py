from __future__ import annotations

import argparse
import shutil
import subprocess
import textwrap
from datetime import datetime
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "mockups" / "out"
SCREENSHOT_DIR = OUT_DIR / "screens"
DEFAULT_DECK = OUT_DIR / "gear5-ui-review-deck.pptx"

ROUTES = [
    ("/", "Role Switch"),
    ("/rider", "Rider Home"),
    ("/search", "Search"),
    ("/rides", "Ride Options"),
    ("/ride-pack", "Ride Pack"),
    ("/myride", "My Ride"),
    ("/myride/details", "My Ride Details"),
    ("/live", "Live Ride"),
    ("/profile", "Profile"),
    ("/driver", "Driver"),
    ("/admin", "Admin"),
    ("/superadmin", "Superadmin"),
]


def find_browsers() -> list[Path]:
    candidates = [
        Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
        Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
        Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe"),
        Path(r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe"),
    ]
    browsers = [candidate for candidate in candidates if candidate.exists()]
    if not browsers:
        raise RuntimeError("Could not find Chrome or Edge.")
    return browsers


def route_to_file(route: str) -> str:
    if route == "/":
        return "root.png"
    return route.strip("/").replace("/", "-") + ".png"


def capture_screenshots(base_url: str, width: int, height: int) -> list[tuple[str, str, Path]]:
    browsers = find_browsers()
    SCREENSHOT_DIR.mkdir(parents=True, exist_ok=True)

    captures: list[tuple[str, str, Path]] = []
    for route, label in ROUTES:
        output = SCREENSHOT_DIR / route_to_file(route)
        url = f"{base_url.rstrip('/')}/#{route}"
        profile_dir = OUT_DIR / "browser-profile"
        last_error: subprocess.CalledProcessError | None = None
        for browser in browsers:
            shutil.rmtree(profile_dir, ignore_errors=True)
            profile_dir.mkdir(parents=True, exist_ok=True)
            command = [
                str(browser),
                "--headless=new",
                "--no-sandbox",
                "--disable-gpu",
                "--disable-gpu-compositing",
                "--disable-dev-shm-usage",
                "--disable-features=VizDisplayCompositor",
                "--hide-scrollbars",
                "--no-first-run",
                f"--user-data-dir={profile_dir}",
                f"--window-size={width},{height}",
                f"--screenshot={output}",
                url,
            ]
            try:
                subprocess.run(command, check=True, cwd=ROOT, capture_output=True, text=True)
                last_error = None
                break
            except subprocess.CalledProcessError as error:
                last_error = error
        if last_error is not None:
            details = (last_error.stderr or last_error.stdout or "").strip()
            raise RuntimeError(f"Could not capture {url}. Browser output: {details}") from last_error
        captures.append((route, label, output))

    shutil.rmtree(OUT_DIR / "browser-profile", ignore_errors=True)
    return captures


def fit_image(slide, image_path: Path, left, top, max_width, max_height):
    with Image.open(image_path) as image:
        width_px, height_px = image.size
    ratio = min(max_width / width_px, max_height / height_px)
    width = int(width_px * ratio)
    height = int(height_px * ratio)
    return slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)


def add_textbox(slide, left, top, width, height, text, size=12, bold=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    return box


def add_note_panel(slide, left, top, width, height):
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    panel.fill.solid()
    panel.fill.background()
    panel.line.fill.background()
    panel.text_frame.clear()
    return panel


def build_deck(captures: list[tuple[str, str, Path]], output: Path, base_url: str):
    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)

    blank_layout = prs.slide_layouts[6]

    title = prs.slides.add_slide(blank_layout)
    add_textbox(title, Inches(0.6), Inches(0.5), Inches(14.8), Inches(0.7), "Gear5 UI Review Deck", 30, True)
    add_textbox(
        title,
        Inches(0.65),
        Inches(1.35),
        Inches(9.6),
        Inches(0.8),
        f"Generated from {base_url} on {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        14,
    )
    instructions = textwrap.dedent(
        """
        How to edit:
        - Move, resize, draw over, or annotate anything directly on the slide.
        - Add arrows/comments for spacing, copy, color, or hierarchy changes.
        - Keep the route label on each slide so Codex knows which screen to update.
        - Save the edited PPTX into mockups/inbox and ask Codex to apply it.

        Tip: if a slide should be very precise, export it as PNG too and place it next to the PPTX.
        """
    ).strip()
    add_textbox(title, Inches(0.7), Inches(2.35), Inches(7.2), Inches(3.2), instructions, 18)
    add_textbox(title, Inches(9.1), Inches(2.35), Inches(5.8), Inches(2.6), "Editable notes area is included on every screen slide.", 22, True)

    for route, label, image_path in captures:
        slide = prs.slides.add_slide(blank_layout)
        add_textbox(slide, Inches(0.45), Inches(0.25), Inches(7.6), Inches(0.45), f"{label}  |  {route}", 18, True)
        fit_image(slide, image_path, Inches(0.55), Inches(0.85), Inches(6.0), Inches(7.85))

        add_textbox(slide, Inches(7.0), Inches(0.95), Inches(8.4), Inches(0.35), "Edit / annotate here", 16, True)
        notes = [
            ("Layout", Inches(7.0), Inches(1.45)),
            ("Spacing", Inches(7.0), Inches(3.1)),
            ("Copy / labels", Inches(7.0), Inches(4.75)),
            ("Behavior notes", Inches(7.0), Inches(6.4)),
        ]
        for heading, left, top in notes:
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(8.25), Inches(1.1))
            box.fill.background()
            box.line.fill.background()
            add_textbox(slide, left + Inches(0.16), top + Inches(0.12), Inches(7.9), Inches(0.32), heading, 11, True)
            add_textbox(slide, left + Inches(0.16), top + Inches(0.48), Inches(7.9), Inches(0.42), "Type notes here or draw directly on the screenshot.", 10)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    return output


def main():
    parser = argparse.ArgumentParser(description="Export Gear5 routes into a UI review PPTX.")
    parser.add_argument("--base-url", default="http://127.0.0.1:5173")
    parser.add_argument("--width", type=int, default=390)
    parser.add_argument("--height", type=int, default=844)
    parser.add_argument("--output", type=Path, default=DEFAULT_DECK)
    args = parser.parse_args()

    captures = capture_screenshots(args.base_url, args.width, args.height)
    deck = build_deck(captures, args.output, args.base_url)
    print(deck)


if __name__ == "__main__":
    main()
